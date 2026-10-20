"""Reading uploads and the upload worker's jobs (plan/slide-uploads-plan.md), with the LLM stubbed. No API calls."""

from pathlib import Path

import docx
import pymupdf
import pytest
from assessment_agent import bank_writing, item_writing
from assessment_agent.bank import KuBankGen, BankMCQ
from curriculum_alignment_agent import deck_kus
from curriculum_alignment_agent.deck_kus import EarlierKu, Link, LinkList, UploadKU, UploadKUList, clean_deck_kus
from ingestion_agent import upload_jobs, uploads
from ingestion_agent.uploads import ScheduleLecture, SyllabusSchedule, UploadError, normalize_syllabus, read_deck, read_syllabus, syllabus_key
from syllabyte_common.models import Slide, SlideDeck, SlideRef

SYLLABUS = [
    "CS-UH 2219E Computational Social Science",
    "Instructor: Prof. Example. Office hours by appointment.",
    "This course introduces causal inference for the social sciences, from potential outcomes to randomized experiments and regression.",
    "Grading: problem sets 40%, midterm 25%, final project 35%.",
    "Week 1: Causal questions and potential outcomes.",
    "Week 2: Selection bias.",
    "Week 3: Randomization.",
]


def make_pdf(path: Path, pages: list[str], footer: str | None = None) -> Path:
    doc = pymupdf.open()
    for i, text in enumerate(pages, start=1):
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 700), text, fontsize=11)
        if footer:
            page.insert_text((50, 800), f"{footer}   {i}", fontsize=8)
    doc.save(path)
    return path


# --------------------------------------------------------------------------
# syllabus
# --------------------------------------------------------------------------


def test_normalizing_ignores_case_spacing_punctuation_and_page_numbers():
    a = "CS-UH 2219E:  Computational Social Science\nPage 1 of 3\nWeek 1 — Causal questions."
    b = "cs uh 2219e computational social science\n\n2\nweek 1 causal questions"
    assert normalize_syllabus(a) == normalize_syllabus(b)
    assert syllabus_key(a) == syllabus_key(b)
    assert syllabus_key(a) != syllabus_key(a + "\nWeek 2: Selection bias.")


def test_normalizing_drops_headers_repeated_on_every_page():
    body = ["Week 1: potential outcomes", "Week 2: selection bias", "Week 3: randomization"]
    with_header = "\n".join(f"Fall 2026 · CS-UH 2219E\n{line}" for line in body)
    assert normalize_syllabus(with_header) == normalize_syllabus("\n".join(body))


def test_a_pdf_and_a_word_copy_of_the_same_syllabus_have_the_same_key(tmp_path):
    pdf = make_pdf(tmp_path / "s.pdf", ["\n".join(SYLLABUS[:4]), "\n".join(SYLLABUS[4:])], footer="Fall 2026 syllabus")
    word = docx.Document()
    for line in SYLLABUS:
        word.add_paragraph(line)
    word.save(tmp_path / "s.docx")
    assert read_syllabus(pdf).key == read_syllabus(tmp_path / "s.docx").key


def test_an_unreadable_syllabus_is_refused_with_a_message(tmp_path):
    with pytest.raises(UploadError, match="PDF or Word"):
        read_syllabus(tmp_path / "s.txt")
    blank = make_pdf(tmp_path / "blank.pdf", [""])
    with pytest.raises(UploadError, match="No text"):
        read_syllabus(blank)


def test_the_lecture_list_is_renumbered_in_order(monkeypatch):
    def fake(**kw):
        return SyllabusSchedule(course_code="CS", course_name="X", lectures=[ScheduleLecture(n=5, title="B", week=2), ScheduleLecture(n=2, title=" A ", week=None)])

    monkeypatch.setattr(uploads, "structured", fake)
    s = uploads.read_schedule("text")
    assert [(l.n, l.title, l.week) for l in s.lectures] == [(1, "A", None), (2, "B", 2)]


def test_the_lecture_list_is_only_read_for_a_new_course(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(uploads, "structured", lambda **kw: calls.append(1) or SyllabusSchedule(course_code="", course_name="X", lectures=[]))
    pdf = make_pdf(tmp_path / "s.pdf", ["\n".join(SYLLABUS)])
    joined = upload_jobs.run_syllabus_job(pdf, is_new=lambda key: False)
    assert "schedule" not in joined and calls == []
    created = upload_jobs.run_syllabus_job(pdf, is_new=lambda key: True)
    assert created["schedule"]["lectures"] == [] and calls == [1]
    assert created["key"] == joined["key"]


# --------------------------------------------------------------------------
# decks
# --------------------------------------------------------------------------


def test_the_same_slides_get_the_same_content_key(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", ["Alpha intro", "Beta details"])
    b = make_pdf(tmp_path / "b.pdf", ["Alpha intro", "Beta details"])
    c = make_pdf(tmp_path / "c.pdf", ["Alpha intro", "Beta details, revised"])
    ra, _ = read_deck(a, tmp_path / "wa", "d1")
    rb, _ = read_deck(b, tmp_path / "wb", "d2")
    rc, _ = read_deck(c, tmp_path / "wc", "d3")
    assert ra.content_key == rb.content_key != rc.content_key
    assert ra.page_count == 2 and ra.first_title == "Alpha intro"


def test_a_read_deck_cites_its_server_id(tmp_path):
    _, deck = read_deck(make_pdf(tmp_path / "a.pdf", ["Alpha"]), tmp_path / "w", "deck-42")
    assert deck.name == "deck-42" and {s.deck_name for s in deck.slides} == {"deck-42"}
    assert deck.has_images and Path(deck.slides[0].image_path).exists()


def test_keynote_and_other_files_are_refused_with_a_message(tmp_path):
    (tmp_path / "a.key").write_bytes(b"x")
    with pytest.raises(UploadError, match="Export To > PDF"):
        uploads.to_pdf(tmp_path / "a.key", tmp_path / "out")
    (tmp_path / "a.png").write_bytes(b"x")
    with pytest.raises(UploadError, match="PDF or PowerPoint"):
        uploads.to_pdf(tmp_path / "a.png", tmp_path / "out")


# --------------------------------------------------------------------------
# KUs of a deck
# --------------------------------------------------------------------------


def uku(id_: str, prereqs: list[str] = (), slides: list[int] = (1,)) -> UploadKU:
    return UploadKU(
        id=id_, title=f"Title {id_}", short_title=id_, summary=f"About {id_}", bloom_level="understand", cognitive_load="low",
        prerequisites=list(prereqs), source_refs=[SlideRef(deck_name="whatever", slide_number=n) for n in slides],
        visual_idea="x", short_worthiness=3, short_worthiness_reason="x",
    )


DECK = SlideDeck(name="d9", source_path="d9.pdf", slides=[Slide(deck_name="d9", slide_number=n, text=f"page {n}") for n in (1, 2, 3)])


def test_clashing_ids_get_a_suffix_and_links_follow_this_decks_kus():
    out = clean_deck_kus([uku("ku-a"), uku("ku-b", ["ku-a"])], DECK, earlier_ids={"ku-a"}, taken_ids={"ku-a"})
    assert [k.id for k in out] == ["ku-a-2", "ku-b"]
    # "ku-a" names this deck's KU (now ku-a-2), not the earlier lecture's
    assert out[1].prerequisites == ["ku-a-2"]


def test_prerequisites_may_name_earlier_lectures_but_nothing_unknown():
    out = clean_deck_kus([uku("ku-b", ["ku-old", "ku-nowhere", "ku-b"])], DECK, earlier_ids={"ku-old"}, taken_ids={"ku-old"})
    assert out[0].prerequisites == ["ku-old"]


def test_cycles_inside_a_deck_are_broken():
    out = clean_deck_kus([uku("ku-a", ["ku-b"]), uku("ku-b", ["ku-a"])], DECK, earlier_ids=set(), taken_ids=set())
    assert out[0].prerequisites == ["ku-b"] and out[1].prerequisites == []


def test_slide_references_point_at_this_deck_and_its_pages():
    out = clean_deck_kus([uku("ku-a", slides=[2, 7])], DECK, earlier_ids=set(), taken_ids=set())
    assert [(r.deck_name, r.slide_number) for r in out[0].source_refs] == [("d9", 2)]


def test_the_link_pass_only_adds_later_to_new_links(monkeypatch):
    new = [EarlierKu(id="ku-n", title="N", summary="", lecture_n=2)]
    later = [EarlierKu(id="ku-l", title="L", summary="", lecture_n=5)]
    monkeypatch.setattr(
        deck_kus, "structured",
        lambda **kw: LinkList(links=[Link(ku_id="ku-l", prereq_id="ku-n"), Link(ku_id="ku-n", prereq_id="ku-l"), Link(ku_id="ku-l", prereq_id="ku-n"), Link(ku_id="ku-x", prereq_id="ku-n")]),
    )
    out = upload_jobs.run_link_pass_job({"course_name": "C", "new_kus": [k.model_dump() for k in new], "later_kus": [k.model_dump() for k in later]})
    assert out == {"links": [{"ku_id": "ku-l", "prereq_id": "ku-n"}]}


# --------------------------------------------------------------------------
# the whole deck job, LLM stubbed
# --------------------------------------------------------------------------


def stub_llm(monkeypatch):
    """KUs: one per page, named after the page's first word, each building on the one before; items: valid MCQs."""

    def kus(**kw):
        text = next(b["text"] for b in kw["content"] if b.get("type") == "text")
        name = text.split("=== DECK '")[1].split("'")[0]
        pages = STUB_PAGES[name]
        out, prev = [], None
        for n, page in enumerate(pages, start=1):
            id_ = f"ku-{page.split()[0].lower()}"
            out.append(uku(id_, [prev] if prev else ([STUB_EARLIER[name]] if name in STUB_EARLIER else []), [n]))
            prev = id_
        return UploadKUList(knowledge_units=out)

    def mcq(i: int) -> BankMCQ:
        return BankMCQ(stem=f"Which is right {i}?", options=["a", "b", "c"], correct_index=0, explanation="Because.")

    def bank(**kw):
        return KuBankGen(basic=mcq(1), applied=mcq(2), advanced=mcq(3), slides=[1], ku_difficulty=0.5)

    def verdicts(**kw):
        return item_writing.Verdicts(verdicts=[])

    monkeypatch.setattr(deck_kus, "structured", kus)
    monkeypatch.setattr(bank_writing, "structured", bank)
    monkeypatch.setattr(item_writing, "structured", verdicts)


STUB_PAGES: dict[str, list[str]] = {}
STUB_EARLIER: dict[str, str] = {}


def test_a_deck_job_returns_kus_and_a_valid_item_bank(tmp_path, monkeypatch):
    stub_llm(monkeypatch)
    STUB_PAGES["d2"] = ["Gamma first", "Delta second"]
    STUB_EARLIER["d2"] = "ku-alpha"
    pdf = make_pdf(tmp_path / "two.pdf", STUB_PAGES["d2"])
    payload = {
        "deck_id": "d2", "course_name": "C", "syllabus_text": "S", "lecture_n": 2, "lecture_title": "Two",
        "earlier": [{"id": "ku-alpha", "title": "Alpha", "summary": "", "lecture_n": 1}], "taken_ids": ["ku-alpha"],
    }
    out = upload_jobs.run_deck_job(payload, pdf, tmp_path / "work", log=lambda m: None)
    assert [(k["id"], k["prerequisites"]) for k in out["kus"]] == [("ku-gamma", ["ku-alpha"]), ("ku-delta", ["ku-gamma"])]
    assert {k["short_title"] for k in out["kus"]} == {"ku-gamma", "ku-delta"}
    assert len(out["items"]) == 6
    assert {i["ku_id"] for i in out["items"]} == {"ku-gamma", "ku-delta"}
    assert {r["deck_name"] for i in out["items"] for r in i["source_refs"]} == {"d2"}
    assert sorted({i["correct_index"] for i in out["items"]}) == [0, 1, 2]  # answers spread over A/B/C
    assert out["cost_usd"] == 0


def test_numbers_in_the_syllabus_text_still_count():
    assert syllabus_key("Course syllabus\nOffered Fall 2025") != syllabus_key("Course syllabus\nOffered Fall 2026")


def make_pptx(path: Path, texts: list[str]) -> Path:
    from pptx import Presentation

    prs = Presentation()
    for t in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = t
    prs.save(path)
    return path


def test_a_worker_without_libreoffice_fails_the_job_instead_of_guessing(tmp_path, monkeypatch):
    monkeypatch.setattr(uploads.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        uploads.to_pdf(make_pptx(tmp_path / "a.pptx", ["Alpha"]), tmp_path / "out")


@pytest.mark.skipif(not (uploads.shutil.which("soffice") or uploads.shutil.which("libreoffice")), reason="LibreOffice is not installed")
def test_a_powerpoint_deck_is_read_as_its_pdf(tmp_path):
    read, deck = read_deck(make_pptx(tmp_path / "a.pptx", ["Alpha", "Beta"]), tmp_path / "w", "d1")
    assert read.pdf_path.endswith(".pdf") and read.page_count == 2 and deck.has_images
    assert read.first_title == "Alpha"
