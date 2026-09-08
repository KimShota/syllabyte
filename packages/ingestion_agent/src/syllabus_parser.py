"""
Syllabus Parser — Component 1 (Ingestion Agent), September scope.

Extracts raw structured text from a course syllabus, whether it's a PDF
or a Word document. This is deliberately "dumb" extraction only: no LLM
calls here. The output (a SyllabusDocument) is what October's LLM-based
Knowledge Unit extraction will consume.

Usage:
    from syllabus_parser import parse_syllabus
    doc = parse_syllabus("path/to/syllabus.pdf")
    print(doc.full_text[:500])
    for section in doc.sections:
        print(section.heading, "->", len(section.body), "chars")
"""

from dataclasses import dataclass, field
from pathlib import Path

import pymupdf  # PyMuPDF (import name 'fitz' is deprecated)
import docx  # python-docx


@dataclass
class SyllabusSection:
    heading: str
    body: str


@dataclass
class SyllabusDocument:
    source_path: str
    full_text: str
    sections: list[SyllabusSection] = field(default_factory=list)


def _parse_pdf(path: Path) -> str:
    text_chunks = []
    with pymupdf.open(path) as pdf:
        for page in pdf:
            text_chunks.append(page.get_text("text"))
    return "\n".join(text_chunks)


def _parse_docx(path: Path) -> str:
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def _naive_section_split(full_text: str) -> list[SyllabusSection]:
    """
    Very naive heuristic splitter: treats short, title-cased or all-caps
    lines as headings and groups subsequent lines under them. This is a
    placeholder — October's LLM-based extraction will replace this with
    real semantic segmentation into Knowledge Units.
    """
    sections: list[SyllabusSection] = []
    current_heading = "Untitled"
    current_body: list[str] = []

    def is_heading(line: str) -> bool:
        stripped = line.strip()
        if not stripped or len(stripped) > 80:
            return False
        return stripped.isupper() or (stripped.istitle() and len(stripped.split()) <= 8)

    for line in full_text.splitlines():
        if is_heading(line):
            if current_body:
                sections.append(SyllabusSection(current_heading, "\n".join(current_body).strip()))
            current_heading = line.strip()
            current_body = []
        else:
            current_body.append(line)

    if current_body:
        sections.append(SyllabusSection(current_heading, "\n".join(current_body).strip()))

    return sections


def parse_syllabus(path: str | Path) -> SyllabusDocument:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        full_text = _parse_pdf(path)
    elif suffix in (".docx", ".doc"):
        full_text = _parse_docx(path)
    else:
        raise ValueError(f"Unsupported syllabus format: {suffix}. Expected .pdf or .docx")

    sections = _naive_section_split(full_text)
    return SyllabusDocument(source_path=str(path), full_text=full_text, sections=sections)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python syllabus_parser.py <path_to_syllabus.pdf|.docx>")
        sys.exit(1)

    doc = parse_syllabus(sys.argv[1])
    print(f"Parsed {len(doc.full_text)} characters into {len(doc.sections)} sections.\n")
    for s in doc.sections[:5]:
        print(f"--- {s.heading} ---")
        print(s.body[:200], "...\n")
